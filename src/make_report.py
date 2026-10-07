from datetime import datetime
import io
import json
from pathlib import Path
import sys
from typing import Annotated

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage
import tyro
from tyro.conf import arg

# STYLES

COLOR_MAP = {
    "TN": {"fill": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"), "font": Font(color="155724", bold=True)},
    "TP": {"fill": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"), "font": Font(color="721C24", bold=True)},
    "FP": {"fill": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"), "font": Font(color="856404", bold=True)},
    "FN": {"fill": PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid"), "font": Font(color="004085", bold=True)},
}

HEADER_FILL = PatternFill(start_color="2A3F54", end_color="2A3F54", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BORDER_THIN = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)


def create_thumbnail_bytes(image_path: Path, crop_middle_third: bool = False, size: tuple[int, int] = (95, 95)) -> io.BytesIO | None:
    """Prepare demo images for table."""
    if not image_path.exists():
        return None
    try:
        with PILImage.open(image_path) as img:
            img = img.convert("RGB")
            if crop_middle_third:
                w, h = img.size
                img = img.crop((int(w * 0.33), 0, int(w * 0.67), h))

            img.thumbnail(size, PILImage.Resampling.LANCZOS)
            output = io.BytesIO()
            img.save(output, format="JPEG", quality=85)
            output.seek(0)
            return output
    except Exception as e:
        print(f"Warning: Failed to process image {image_path}: {e}", file=sys.stderr)
        return None


def build_metadata_sheet(ws, meta: dict):
    """Sheet 1 (metrics & notes)"""
    ws.title = "Summary & Metadata"
    ws.views.sheetView[0].showGridLines = True

    # title
    ws["A1"] = "SUMMARY & METRICS"
    ws["A1"].font = Font(size=14, bold=True, color="2A3F54")
    ws.merge_cells("A1:C1")

    # parameters
    cm = meta.get("confusion_matrix", {})
    m = meta.get("metrics", {})
    perf = meta.get("performance", {})
    sc = meta.get("score_stats", {})

    items = [
        ("Run Timestamp", meta.get("timestamp", "N/A")),
        ("Model Path", meta.get("model_path", "N/A")),
        ("Test Dataset", meta.get("test_dataset", "N/A")),
        ("Threshold Used", meta.get("threshold_used", "N/A")),
        ("Inference Device", meta.get("device", "N/A")),
        ("---", "---"),
        ("Total Images", meta.get("total_processed", 0)),
        ("True Positives (TP)", cm.get("TP", 0)),
        ("True Negatives (TN)", cm.get("TN", 0)),
        ("False Positives (FP)", cm.get("FP", 0)),
        ("False Negatives (FN)", cm.get("FN", 0)),
        ("---", "---"),
        ("Accuracy", f"{m.get('accuracy', 0.0):.2%}"),
        ("Recall", f"{m.get('recall', 0.0):.2%}"),
        ("Precision", f"{m.get('precision', 0.0):.2%}"),
        ("F1-Score", f"{m.get('f1', 0.0):.2%}"),
        ("---", "---"),
        ("Avg Inference Time", f"{perf.get('avg_inference_time_sec', 0.0):.4f} s"),
        ("FPS", f"{perf.get('fps', 0.0):.1f}"),
        ("---", "---"),
        ("Normal Scores (min/mean/max)", f"{sc.get('normal', {}).get('min')} / {sc.get('normal', {}).get('mean')} / {sc.get('normal', {}).get('max')}"),
        ("Anomalous Scores (min/mean/max)", f"{sc.get('anomalous', {}).get('min')} / {sc.get('anomalous', {}).get('mean')} / {sc.get('anomalous', {}).get('max')}"),
    ]

    for row_idx, (k, v) in enumerate(items, start=3):
        if k == "---":
            continue
        cell_k = ws[f"A{row_idx}"]
        cell_v = ws[f"B{row_idx}"]

        cell_k.value = k
        cell_k.font = Font(bold=True, color="4F4F4F")
        cell_k.fill = PatternFill(start_color="F7F9FA", end_color="F7F9FA", fill_type="solid")

        cell_v.value = v
        cell_v.font = Font(bold=False)
        cell_v.alignment = Alignment(horizontal="left", vertical="center")

        cell_k.border = BORDER_THIN
        cell_v.border = BORDER_THIN

    # notes
    ws["D2"] = "NOTES"
    ws["D2"].font = HEADER_FONT
    ws["D2"].fill = HEADER_FILL
    ws["D2"].alignment = Alignment(horizontal="center", vertical="center")

    notes_box = ws["D3"]
    notes_box.alignment = Alignment(vertical="top", wrap_text=True)
    notes_box.border = Border(
        left=Side(style="medium", color="2A3F54"),
        right=Side(style="medium", color="2A3F54"),
        top=Side(style="medium", color="2A3F54"),
        bottom=Side(style="medium", color="2A3F54"),
    )
    notes_box.fill = PatternFill(start_color="FFFDF8", end_color="FFFDF8", fill_type="solid")
    ws.merge_cells("D3:E20")

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 35


def build_results_sheet(ws, records: list[dict]):
    """Sheet 2 (results)"""
    ws.title = "Detailed Results"
    ws.views.sheetView[0].showGridLines = True

    headers = ["#", "Filename", "Ground Truth", "Prediction", "Status", "Score", "Photo", "Heatmap"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    image_references = []

    for row_idx, r in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 75

        res_type = r.get("result_type", "")
        style = COLOR_MAP.get(res_type, {})

        ws.cell(row=row_idx, column=1, value=row_idx - 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=2, value=r.get("filename", "")).alignment = Alignment(vertical="center")
        ws.cell(row=row_idx, column=3, value=r.get("gt", "")).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=4, value=r.get("predict", "")).alignment = Alignment(horizontal="center", vertical="center")

        # colored status cell
        status_cell = ws.cell(row=row_idx, column=5, value=res_type)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        if style:
            status_cell.fill = style["fill"]
            status_cell.font = style["font"]

        score_val = r.get("score")
        score_cell = ws.cell(row=row_idx, column=6, value=score_val if score_val is not None else "N/A")
        score_cell.alignment = Alignment(horizontal="center", vertical="center")

        # orig photo preview
        photo_p = Path(r["photo_path"]) if r.get("photo_path") else None
        if photo_p and photo_p.exists():
            thumb_bio = create_thumbnail_bytes(photo_p, crop_middle_third=False)
            if thumb_bio:
                image_references.append(thumb_bio)
                img_obj = OpenpyxlImage(thumb_bio)
                ws.add_image(img_obj, f"G{row_idx}")

        # heatmap preview
        vis_p = Path(r["vis_path"]) if r.get("vis_path") else None
        if vis_p and vis_p.exists():
            heat_bio = create_thumbnail_bytes(vis_p, crop_middle_third=True)
            if heat_bio:
                image_references.append(heat_bio)
                img_obj = OpenpyxlImage(heat_bio)
                ws.add_image(img_obj, f"H{row_idx}")

        for col_idx in range(1, 9):
            ws.cell(row=row_idx, column=col_idx).border = BORDER_THIN

    col_widths = {"A": 6, "B": 28, "C": 15, "D": 15, "E": 12, "F": 12, "G": 16, "H": 16}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # filters
    ws.auto_filter.ref = f"A1:H{len(records) + 1}"
    return image_references


def main(
    eval_dir: Annotated[Path | None, arg(aliases=["-d"])] = None,
    metadata: Annotated[Path | None, arg(aliases=["-m"])] = None,
    results: Annotated[Path | None, arg(aliases=["-r"])] = None,
    output: Annotated[Path | None, arg(aliases=["-o"])] = None,
):
    """Generate an Excel evaluation report with thumbnails and status color highlighting.

    Args:
        eval_dir: Path to eval run directory containing both metadata.json and results.json.
        metadata: Explicit path to metadata.json.
        results: Explicit path to results.json.
        output: Destination path for .xlsx report. If None, saves to reports/report_{timestamp}.xlsx.
    """
    # search for source files
    if eval_dir:
        metadata = eval_dir / "metadata.json"
        results = eval_dir / "results.json"

    if not metadata or not metadata.exists():
        print(f"Error: metadata.json not found ({metadata}). Provide -d or -m.", file=sys.stderr)
        sys.exit(1)

    if not results or not results.exists():
        print(f"Error: results.json not found ({results}). Provide -d or -r.", file=sys.stderr)
        sys.exit(1)

    # output file
    if output is None:
        reports_dir = Path("reports")
        reports_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output = reports_dir / f"report_{timestamp}.xlsx"
    else:
        output.parent.mkdir(parents=True, exist_ok=True)

    print(f">>> Reading inputs:\n  Metadata: {metadata}\n  Results : {results}")

    with open(metadata, "r", encoding="utf-8") as f:
        meta_data = json.load(f)

    with open(results, "r", encoding="utf-8") as f:
        records_data = json.load(f)

    # build excel book
    wb = Workbook()
    ws_meta = wb.active
    build_metadata_sheet(ws_meta, meta_data)

    ws_results = wb.create_sheet()
    _keep_alive_imgs = build_results_sheet(ws_results, records_data)

    print(f">>> Saving Excel workbook to: {output}...")
    wb.save(output)
    print(f">>> Done! Report generated successfully.")


if __name__ == "__main__":
    tyro.cli(main)