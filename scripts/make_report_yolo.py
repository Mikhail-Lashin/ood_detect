from datetime import datetime
import io
import json
from pathlib import Path
import sys
from typing import Annotated

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage
import tyro
from tyro.conf import arg

# STYLES

COLOR_MAP = {
    "CORRECT": {"fill": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"), "font": Font(color="155724", bold=True)},
    "ERROR": {"fill": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"), "font": Font(color="721C24", bold=True)},
}

HEADER_FILL = PatternFill(start_color="2A3F54", end_color="2A3F54", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BORDER_THIN = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)


def create_thumbnail_bytes(image_path: Path, size: tuple[int, int] = (95, 95)) -> io.BytesIO | None:
    """Prepare demo images for table."""
    if not image_path.exists():
        return None
    try:
        with PILImage.open(image_path) as img:
            img = img.convert("RGB")
            img.thumbnail(size, PILImage.Resampling.LANCZOS)
            output = io.BytesIO()
            img.save(output, format="JPEG", quality=85)
            output.seek(0)
            return output
    except Exception as e:
        print(f"Warning: Failed to process image {image_path}: {e}", file=sys.stderr)
        return None


def build_metadata_sheet(ws, meta: dict):
    """Sheet 1 (metrics & notes)"""
    ws.title = "Summary & Metadata"
    ws.views.sheetView[0].showGridLines = True

    # title
    ws["A1"] = "YOLO CLASSIFICATION SUMMARY & METRICS"
    ws["A1"].font = Font(size=14, bold=True, color="2A3F54")
    ws.merge_cells("A1:C1")

    # parameters
    perf = meta.get("performance", {})
    classes = meta.get("classes", [])
    cm = meta.get("confusion_matrix", [])
    pcm = meta.get("per_class_metrics", {})

    items = [
        ("Run Timestamp", meta.get("timestamp", "N/A")),
        ("Model Path", meta.get("model_path", "N/A")),
        ("Test Dataset", meta.get("test_dataset", "N/A")),
        ("Inference Device", meta.get("device", "N/A")),
        ("---", "---"),
        ("Total Images", meta.get("total_processed", 0)),
        ("Overall Accuracy", f"{meta.get('accuracy', 0.0):.2%}"),
        ("---", "---"),
        ("Avg Inference Time", f"{perf.get('avg_inference_time_sec', 0.0):.4f} s"),
        ("FPS", f"{perf.get('fps', 0.0):.1f}"),
    ]

    curr_row = 3
    for k, v in items:
        if k == "---":
            continue
        cell_k = ws[f"A{curr_row}"]
        cell_v = ws[f"B{curr_row}"]

        cell_k.value = k
        cell_k.font = Font(bold=True, color="4F4F4F")
        cell_k.fill = PatternFill(start_color="F7F9FA", end_color="F7F9FA", fill_type="solid")

        cell_v.value = v
        cell_v.font = Font(bold=False)
        cell_v.alignment = Alignment(horizontal="left", vertical="center")

        cell_k.border = BORDER_THIN
        cell_v.border = BORDER_THIN
        curr_row += 1

    # confusion matrix
    if classes and cm:
        curr_row += 1
        ws.cell(row=curr_row, column=1, value="CONFUSION MATRIX").font = Font(bold=True, size=11, color="2A3F54")
        curr_row += 1

        headers = ["True \\ Pred"] + classes
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=c_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, cls_name in enumerate(classes):
            row_num = curr_row + 1 + r_idx
            label_cell = ws.cell(row=row_num, column=1, value=cls_name)
            label_cell.font = Font(bold=True)
            label_cell.border = BORDER_THIN

            for c_idx in range(len(classes)):
                val_cell = ws.cell(row=row_num, column=2 + c_idx, value=cm[r_idx][c_idx])
                val_cell.border = BORDER_THIN
                val_cell.alignment = Alignment(horizontal="center", vertical="center")
                if r_idx == c_idx:
                    val_cell.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

        curr_row = curr_row + len(classes) + 1

    # per-class metrics
    if pcm:
        curr_row += 1
        ws.cell(row=curr_row, column=1, value="PER-CLASS METRICS").font = Font(bold=True, size=11, color="2A3F54")
        curr_row += 1

        m_headers = ["Class", "Recall", "Precision", "F1-Score", "Samples"]
        for c_idx, h in enumerate(m_headers, start=1):
            cell = ws.cell(row=curr_row, column=c_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, (cls_name, m_data) in enumerate(pcm.items()):
            row_num = curr_row + 1 + r_idx
            ws.cell(row=row_num, column=1, value=cls_name).font = Font(bold=True)
            ws.cell(row=row_num, column=2, value=f"{m_data.get('recall', 0.0):.2%}").alignment = Alignment(horizontal="right")
            ws.cell(row=row_num, column=3, value=f"{m_data.get('precision', 0.0):.2%}").alignment = Alignment(horizontal="right")
            ws.cell(row=row_num, column=4, value=f"{m_data.get('f1', 0.0):.2%}").alignment = Alignment(horizontal="right")
            ws.cell(row=row_num, column=5, value=m_data.get("samples", 0)).alignment = Alignment(horizontal="right")

            for c_idx in range(1, 6):
                ws.cell(row=row_num, column=c_idx).border = BORDER_THIN

    # notes
    '''ws["D2"] = "NOTES"
    ws["D2"].font = HEADER_FONT
    ws["D2"].fill = HEADER_FILL
    ws["D2"].alignment = Alignment(horizontal="center", vertical="center")

    notes_box = ws["D3"]
    notes_box.alignment = Alignment(vertical="top", wrap_text=True)
    notes_box.border = Border(
        left=Side(style="medium", color="2A3F54"),
        right=Side(style="medium", color="2A3F54"),
        top=Side(style="medium", color="2A3F54"),
        bottom=Side(style="medium", color="2A3F54"),
    )
    notes_box.fill = PatternFill(start_color="FFFDF8", end_color="FFFDF8", fill_type="solid")
    ws.merge_cells("D3:E20")'''

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 35


def build_results_sheet(ws, records: list[dict]):
    """Sheet 2 (results)"""
    ws.title = "Detailed Results"
    ws.views.sheetView[0].showGridLines = True

    headers = ["#", "Filename", "Ground Truth", "Prediction", "Confidence", "Status", "Photo"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    image_references = []

    for row_idx, r in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 75

        res_type = r.get("result_type", "CORRECT")
        style = COLOR_MAP.get(res_type, {})

        ws.cell(row=row_idx, column=1, value=row_idx - 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=2, value=r.get("filename", "")).alignment = Alignment(vertical="center")
        ws.cell(row=row_idx, column=3, value=r.get("gt", "")).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=4, value=r.get("predict", "")).alignment = Alignment(horizontal="center", vertical="center")

        conf_val = r.get("confidence")
        conf_str = f"{conf_val:.2%}" if isinstance(conf_val, (float, int)) else "N/A"
        ws.cell(row=row_idx, column=5, value=conf_str).alignment = Alignment(horizontal="center", vertical="center")

        # colored status cell
        status_cell = ws.cell(row=row_idx, column=6, value=res_type)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        if style:
            status_cell.fill = style["fill"]
            status_cell.font = style["font"]

        # orig photo preview
        photo_p = Path(r["photo_path"]) if r.get("photo_path") else None
        if photo_p and photo_p.exists():
            thumb_bio = create_thumbnail_bytes(photo_p)
            if thumb_bio:
                image_references.append(thumb_bio)
                img_obj = OpenpyxlImage(thumb_bio)
                ws.add_image(img_obj, f"G{row_idx}")

        for col_idx in range(1, 8):
            ws.cell(row=row_idx, column=col_idx).border = BORDER_THIN

    col_widths = {"A": 6, "B": 28, "C": 18, "D": 18, "E": 14, "F": 14, "G": 16}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # filters
    ws.auto_filter.ref = f"A1:G{len(records) + 1}"
    return image_references


def main(
    eval_dir: Annotated[Path | None, arg(aliases=["-d"])] = None,
    metadata: Annotated[Path | None, arg(aliases=["-m"])] = None,
    results: Annotated[Path | None, arg(aliases=["-r"])] = None,
    output: Annotated[Path | None, arg(aliases=["-o"])] = None,
):
    """Generate an Excel evaluation report for YOLO classification.

    Args:
        eval_dir: Path to directory containing both metadata.json and results.json.
        metadata: Explicit path to metadata.json.
        results: Explicit path to results.json.
        output: Destination path for .xlsx report. If None, saves to reports/yolo_report_{timestamp}.xlsx.
    """
    # search for source files
    if eval_dir:
        metadata = eval_dir / "metadata.json"
        results = eval_dir / "results.json"

    if not metadata or not metadata.exists():
        print(f"Error: metadata.json not found ({metadata}). Provide -d or -m.", file=sys.stderr)
        sys.exit(1)

    if not results or not results.exists():
        print(f"Error: results.json not found ({results}). Provide -d or -r.", file=sys.stderr)
        sys.exit(1)

    # output file
    if output is None:
        reports_dir = Path("reports")
        reports_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output = reports_dir / f"yolo_report_{timestamp}.xlsx"
    else:
        output.parent.mkdir(parents=True, exist_ok=True)

    print(f">>> Reading inputs:\n  Metadata: {metadata}\n  Results : {results}")

    with open(metadata, "r", encoding="utf-8") as f:
        meta_data = json.load(f)

    with open(results, "r", encoding="utf-8") as f:
        records_data = json.load(f)

    # build excel book
    wb = Workbook()
    ws_meta = wb.active
    build_metadata_sheet(ws_meta, meta_data)

    ws_results = wb.create_sheet()
    _keep_alive_imgs = build_results_sheet(ws_results, records_data)

    print(f">>> Saving Excel workbook to: {output}...")
    wb.save(output)
    print(f">>> Done! YOLO report generated successfully.")


if __name__ == "__main__":
    tyro.cli(main)